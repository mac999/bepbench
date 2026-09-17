"""Word (.docx) export.

BEPs are still reviewed, redlined and signed in Word, so the export has to be a
real Word document — styled headings, native tables, a cover page and a table of
contents field — not an HTML file renamed.

Per the house rule on document metadata, the author/company fields are set from
the plan itself and never left carrying a tool's default.
"""

from __future__ import annotations

import re
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from ..models import Project
from ..schema.models import Field, Framework
from ..scoring import ScoreReport
from ..scoring.rules import table_rows

try:
    import docx
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover - optional dependency
    DOCX_AVAILABLE = False


class DocxUnavailable(RuntimeError):
    pass


ACCENT = "1D6FA5"


def _shade(cell, colour: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), colour)
    cell._tc.get_or_add_tcPr().append(shading)


def _toc_field(paragraph) -> None:
    """Insert a real TOC field so Word can refresh page numbers itself."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = r'TOC \o "1-2" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose “Update Field” to build the contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, placeholder, end):
        run._r.append(node)


def _field_paragraphs(document, field: Field, value: Any) -> None:
    if value in (None, "", []):
        paragraph = document.add_paragraph("Not yet completed.")
        paragraph.runs[0].italic = True
        paragraph.runs[0].font.color.rgb = RGBColor(0x90, 0x94, 0x9B)
        return

    if field.type == "table":
        rows = table_rows(value)
        if not rows:
            document.add_paragraph("Not yet completed.")
            return
        table = document.add_table(rows=1, cols=len(field.columns))
        table.style = "Table Grid"
        header = table.rows[0]
        for cell, column in zip(header.cells, field.columns):
            cell.text = column.label
            _shade(cell, "F1F3F6")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(8.5)
        for row in rows:
            cells = table.add_row().cells
            for cell, column in zip(cells, field.columns):
                cell.text = str(row.get(column.id, "") or "—")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
        document.add_paragraph()
        return

    if field.type in ("multiselect", "list"):
        for item in (value if isinstance(value, list) else [value]):
            if str(item).strip():
                document.add_paragraph(str(item), style="List Bullet")
        return

    if field.type == "boolean":
        document.add_paragraph("Yes" if value else "No")
        return

    for block in str(value).split("\n\n"):
        if block.strip():
            document.add_paragraph(block.strip())


def build(project: Project, framework: Framework, report: ScoreReport,
          *, include_score: bool = True) -> bytes:
    if not DOCX_AVAILABLE:
        raise DocxUnavailable(
            "Word export needs the optional 'python-docx' package (pip install python-docx)."
        )

    document = docx.Document()

    # Document properties: the plan's own identity, never a tool default.
    core = document.core_properties
    core.title = f"BIM Execution Plan — {project.name}"
    core.subject = framework.name
    core.author = project.owner or project.client or ""
    core.company = project.client or ""
    core.comments = ""
    core.category = "BIM Execution Plan"

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # --- cover ---------------------------------------------------------
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("BIM Execution Plan")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string(ACCENT)

    subtitle = document.add_paragraph()
    run = subtitle.add_run(project.name)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x44, 0x48, 0x50)

    facts = [
        ("Framework", f"{framework.name} (v{framework.version})"),
        ("Appointing party", project.client or "—"),
        ("Project reference", project.reference or "—"),
        ("Information manager", project.owner or "—"),
        ("Status", project.status),
        ("Generated", report.generated_at.replace("T", " ").replace("+00:00", " UTC")),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in facts:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
        _shade(cells[0], "F8F9FB")
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    if include_score:
        document.add_paragraph()
        heading = document.add_paragraph()
        run = heading.add_run(f"Completeness score  {report.score:.0f}/100   ·   {report.maturity_label}")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string(ACCENT)
        counts = report.issue_counts()
        document.add_paragraph(
            f"Template coverage {report.coverage:.0f}% · required fields answered "
            f"{report.required_coverage:.0f}% · {counts['blocker']} blocking and "
            f"{counts['major']} major issues open · "
            f"{'ready to issue' if report.is_ready else 'not ready to issue'}."
        )

    document.add_section(WD_SECTION.NEW_PAGE)

    # --- contents ------------------------------------------------------
    document.add_heading("Contents", level=1)
    _toc_field(document.add_paragraph())
    document.add_section(WD_SECTION.NEW_PAGE)

    # --- assessment ----------------------------------------------------
    if include_score:
        document.add_heading("Completeness Assessment", level=1)
        document.add_paragraph(report.maturity_note)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells,
                               ["Section", "Score", "Coverage", "Required answered"]):
            cell.text = label
            _shade(cell, "F1F3F6")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(8.5)
        for section in report.sections:
            cells = table.add_row().cells
            cells[0].text = section.title
            cells[1].text = f"{section.score:.0f}"
            cells[2].text = f"{section.coverage:.0f}%"
            cells[3].text = f"{section.required_complete}/{section.required_total}"

        if report.issues:
            document.add_heading("Outstanding issues", level=2)
            for issue in report.issues:
                paragraph = document.add_paragraph(style="List Bullet")
                run = paragraph.add_run(f"[{issue.severity}] ")
                run.bold = True
                paragraph.add_run(f"{issue.section_title or issue.section_id}: {issue.message}")
                if issue.detail:
                    detail = document.add_paragraph(issue.detail)
                    detail.paragraph_format.left_indent = Pt(24)
                    detail.runs[0].italic = True
                    detail.runs[0].font.size = Pt(9)
        document.add_section(WD_SECTION.NEW_PAGE)

    # --- body ----------------------------------------------------------
    values = project.values()
    for index, section in enumerate(framework.sections, start=1):
        document.add_heading(f"{index}. {section.title}", level=1)
        if section.reference:
            paragraph = document.add_paragraph(f"Reference: {section.reference}")
            paragraph.runs[0].italic = True
            paragraph.runs[0].font.size = Pt(9)
            paragraph.runs[0].font.color.rgb = RGBColor(0x70, 0x74, 0x7C)
        if section.intent:
            paragraph = document.add_paragraph(section.intent)
            paragraph.runs[0].italic = True

        section_values = values.get(section.id) or {}
        for field in section.fields:
            document.add_heading(field.label + (" *" if field.required else ""), level=2)
            _field_paragraphs(document, field, section_values.get(field.id))

    closing = document.add_paragraph(
        f"Fields completed {report.fields_filled}/{report.fields_total}. "
        "Fields marked * are required by the framework."
    )
    closing.runs[0].font.size = Pt(8.5)
    closing.runs[0].font.color.rgb = RGBColor(0x70, 0x74, 0x7C)

    buffer = BytesIO()
    document.save(buffer)
    return _scrub_metadata(buffer.getvalue(), project)


# ---------------------------------------------------------------------------
# Metadata hygiene
# ---------------------------------------------------------------------------

def _scrub_metadata(raw: bytes, project: Project) -> bytes:
    """Strip the authoring template's fingerprints from the saved file.

    python-docx builds on a stock Word template, which leaves behind an
    ``Application`` of "Microsoft Macintosh Word", a 2013 creation date and a
    thumbnail image from the template. None of that is true of this document, so
    it is replaced with the plan's own identity or removed.
    """
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    source = zipfile.ZipFile(BytesIO(raw))
    out_buffer = BytesIO()

    with zipfile.ZipFile(out_buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            name = item.filename
            if name == "docProps/thumbnail.jpeg":
                continue                      # a picture of someone else's document
            data = source.read(name)

            if name == "docProps/app.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<Application>.*?</Application>", "<Application></Application>", text)
                text = re.sub(r"<Company/>|<Company>.*?</Company>",
                              f"<Company>{_xml_escape(project.client)}</Company>", text)
                text = re.sub(r"<Manager/>|<Manager>.*?</Manager>",
                              f"<Manager>{_xml_escape(project.owner)}</Manager>", text)
                text = re.sub(r"<Template>.*?</Template>", "<Template></Template>", text)
                data = text.encode("utf-8")

            elif name == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r'(<dcterms:created[^>]*>).*?(</dcterms:created>)',
                              rf"\g<1>{now}\g<2>", text)
                text = re.sub(r'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)',
                              rf"\g<1>{now}\g<2>", text)
                data = text.encode("utf-8")

            elif name == "_rels/.rels":
                text = data.decode("utf-8")
                text = re.sub(r"<Relationship[^>]*thumbnail[^>]*/>", "", text)
                data = text.encode("utf-8")

            out.writestr(item, data)

    return out_buffer.getvalue()


def _xml_escape(value: str) -> str:
    return (str(value or "").replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))
